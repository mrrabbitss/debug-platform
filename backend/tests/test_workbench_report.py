"""Focused 0.4.0 report tests. Synthetic receipts and isolated SQLite only."""
from copy import deepcopy
import hashlib
from html.parser import HTMLParser
from pathlib import Path
import re
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader
from reportlab.pdfbase import pdfmetrics
import pytest
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker

from app.core.db import Base
from app.core.utils import json_dumps, sha256_file
from app.models import AnalysisRun, Case, KnowledgeDocument, Report
from app.services import report, workbench
from app.services.category_report import markdown_blocks, markdown_html, report_markdown, _coverage
from app.services.diagnosis_contract import validate_llm_diagnosis
from app.services.report_contract import (
    CHAPTERS, LAYERS, builtin_report_template, report_instructions,
    resolve_report_snapshot, table_cells, validate_report_markdown,
)
from app.services.storage import StorageService
from app.workbench_models import WorkbenchRecord


IDS = {"EVT-synthetic-1", "EVT-synthetic-2"}


def template(content, key="DOC-report-guide", category="connection"):
    return {"id": key, "category": category, "content": content, "version": 1,
            "sha256": hashlib.sha256(content.encode()).hexdigest()}


def sample_markdown():
    text = "# 合成组网诊断报告\n\n## 一、组网总览\n\n"
    text += "| 设备 | 型号 | MAC | IP | 上行方式 | 级联关系 | SmartLink版本 | ApOnlineFlag | SyncStatus | 组网状态 | 证据 |\n"
    text += "|---|---|---|---|---|---|---|---|---|---|---|\n"
    for index in (1, 2):
        text += f"| AP{index} | 待确认 | 待确认 | 待确认 | LAN | 待确认 | 待确认 | 0 | 待确认 | 该时刻离线 | [[EVT-synthetic-{index}]] |\n"
    text += "\n## 二、异常设备时间轨迹\n\n"
    for index in (1, 2):
        text += f"### AP{index} — 离线\n\n| 时间 | 事件 | 关键日志/来源 | 证据引用 |\n|---|---|---|---|\n"
        text += f"| 2026-01-01 10:00:0{index} | **链路中断** | `carrier lost` | [[EVT-synthetic-{index}]] |\n\n"
    text += "## 三、异常AP分析推理\n\n"
    for index in (1, 2):
        text += f"### AP{index} — 离线\n\n#### {LAYERS[0]}\n\n"
        text += f"- 记录到链路中断。\n\n> 证据：`carrier lost` [[EVT-synthetic-{index}]]\n\n结论：物理链路异常 ❌。\n\n"
        for layer in LAYERS[1:]:
            text += f"#### {layer}\n\n结论：未继续排查：下层异常。\n\n"
        text += f"因果链：链路中断导致设备在该时刻离线，链路中断的更深层原因待确认。[[EVT-synthetic-{index}]]\n\n"
    text += "## 四、结论与后续方向\n\n### 根因总览\n\n"
    text += "| AP | 故障现象 | 根因 | 推理结论 | 置信度 | 后续方向 |\n|---|---|---|---|---|---|\n"
    for index in (1, 2):
        text += f"| AP{index} | 离线 | 链路中断，深层原因待确认 | 中断导致离线 [[EVT-synthetic-{index}]] | P3（可能） | 采集端口状态和对端日志 |\n"
    text += "\n1. 处理优先级：P0；核对 AP1 对端端口状态。\n2. 处理优先级：P1；核对 AP2 对端端口状态。\n\n"
    text += "### 故障树覆盖\n\n待确认：本次未提供适用账本，总数未知。\n\n### 缺失证据\n\n"
    text += "| 缺失项 | 影响 | 获取方式 |\n|---|---|---|\n| 对端端口记录 | 深层原因待确认 | 采集对应对端日志并核对时间 |\n"
    return text


def payload(markdown=""):
    return {"report_markdown": markdown, "summary": "合成链路中断，深层原因待确认。",
            "confirmed_facts": [{"statement": "记录到链路中断", "evidence_ids": ["EVT-synthetic-1"]}],
            "hypotheses": [{"rank": 1, "title": "链路中断", "description": "链路中断的更深层原因待确认。",
                            "supporting_evidence": ["EVT-synthetic-1"], "confidence_score": 0.6,
                            "confidence_level": "MEDIUM", "priority": "P0"}],
            "recommended_actions": [{"priority": "P0", "action": "采集对端端口日志", "reason": "定位断点", "expected_result": "取得对应时间记录"}],
            "missing_information": ["对端端口记录"], "suspected_modules": [], "limitations": ["仅为合成资料，未确认深层原因。"]}


def test_valid_four_chapters_and_both_template_keywords():
    snapshot = builtin_report_template()
    validate_report_markdown(sample_markdown(), IDS, snapshot)
    for name in ("template_snapshot", "report_template"):
        value = validate_llm_diagnosis(payload(sample_markdown()), IDS, case_evidence_ids=IDS, **{name: snapshot})
        assert value["report_markdown"] == sample_markdown()
        assert value["hypotheses"][0]["priority"] == "P0"


@pytest.mark.parametrize("old,new", [
    ("## 一、组网总览", "## 一、无关章节"),
    ("## 四、结论与后续方向", "### 四、结论与后续方向"),
    ("### AP1 — 离线", "### AP1与AP2 — 离线"),
    ("#### 第一层：物理链路检查", "#### 第三层：WiFi业务层检查"),
    ("因果链：", "描述："),
    ("P3（可能）", "HIGH"),
    ("P3（可能）", "P0"),
    ("结论：未继续排查：下层异常。", "结论：UDM通信正常。[[EVT-synthetic-1]]"),
    ("[[EVT-synthetic-1]]", "[[EVT-another-case]]"),
    ("[[EVT-synthetic-1]]", "[[DOC-method-only]]"),
    ("[[EVT-synthetic-1]]", "[[bad ref]]"),
    ("[[EVT-synthetic-1]]", "`[[EVT-synthetic-1]]`"),
    ("[[EVT-synthetic-1]]", "[link [[EVT-synthetic-1]]](javascript:bad)"),
    ("[[EVT-synthetic-1]]", r"\[[EVT-synthetic-1]]"),
    ("[[EVT-synthetic-1]]", "gw.log:L1"),
])
def test_invalid_structure_and_receipts_rejected(old, new):
    with pytest.raises(ValueError):
        validate_report_markdown(sample_markdown().replace(old, new), IDS, builtin_report_template())


def test_hidden_unknown_reference_and_fake_code_headings_rejected():
    with pytest.raises(ValueError, match="unknown"):
        validate_report_markdown(sample_markdown() + "\n```\n[[EVT-not-this-case]]\n```", IDS, builtin_report_template())
    with pytest.raises(ValueError, match="chapter"):
        validate_report_markdown("```\n" + sample_markdown() + "\n```\n[[EVT-synthetic-1]]", IDS, builtin_report_template())


def test_method_id_is_not_report_fact_even_if_in_general_allowlist():
    with pytest.raises(ValueError, match="unknown"):
        validate_llm_diagnosis(payload(sample_markdown().replace("EVT-synthetic-1", "DOC-guide")),
                               IDS | {"DOC-guide"}, report_template=builtin_report_template(), case_evidence_ids=IDS)


@pytest.mark.parametrize("content", ["请描述连接问题，说明证据和缺口。", "# 报告规范\n\n## 目的\n说明因果。\n\n## 格式\n引用证据。"])
def test_published_prose_specification_is_guidance_not_a_scaffold(content):
    validate_report_markdown("# 连接诊断\n\n本次记录到链路中断 [[EVT-synthetic-1]]", IDS, template(content))


def test_legacy_markdown_not_forced_into_new_chapters():
    validate_llm_diagnosis(payload("# 历史报告\n旧段落"), IDS)
    v1 = builtin_report_template()
    v1["version"] = 1
    validate_report_markdown("旧格式段落", IDS, v1)


def test_snapshot_is_copied_and_hash_mismatch_does_not_fall_back():
    original = {"problem_category": "unknown", "report_template": builtin_report_template()}
    instructions = report_instructions(original)
    instructions["report_template"]["content"] = "tampered"
    assert original["report_template"]["version"] == 2
    assert instructions["report_template"]["content"] != original["report_template"]["content"]
    with pytest.raises(ValueError, match="hash"):
        resolve_report_snapshot(instructions)


def test_category_suggestion_accepts_admin_id_and_requires_receipts():
    data = payload()
    data.update(suggested_problem_category="WB-SYNTHETIC-category", category_reason="观察支持该类别 [[EVT-synthetic-1]]")
    assert validate_llm_diagnosis(data, IDS, case_evidence_ids=IDS)["suggested_problem_category"] == "WB-SYNTHETIC-category"
    data["category_reason"] = "只有方法 [[DOC-guide]]"
    with pytest.raises(ValueError, match="Category suggestion"):
        validate_llm_diagnosis(data, IDS | {"DOC-guide"}, case_evidence_ids=IDS)
    data.pop("category_reason")
    with pytest.raises(ValueError, match="together"):
        validate_llm_diagnosis(data, IDS)


@pytest.fixture
def isolated(tmp_path, monkeypatch):
    engine = create_engine(f"sqlite:///{tmp_path / 'report-tests.db'}")
    Base.metadata.create_all(engine)
    factory = sessionmaker(engine, expire_on_commit=False)
    managed = StorageService(tmp_path / "storage")
    monkeypatch.setattr(report, "SessionLocal", factory)
    monkeypatch.setattr(report, "storage", managed)
    with factory() as db:
        case = Case(id="CASE-synthetic-report", title="合成组网离线检查", device_type="AP", problem_category="network")
        db.add(case)
        db.flush()
        config = workbench.capture_configuration(db, case)
        evidence = [{"evidence_id": f"EVT-synthetic-{index}", "source_type": "log_event", "source_file": f"合成/ap{index}.log",
                     "line_start": 10, "line_end": 12, "content": f"AP{index} carrier lost", "timestamp": f"2026-01-01 10:00:0{index}"} for index in (1, 2)]
        evidence.append({"evidence_id": "DOC-guide", "source_type": "fault_tree", "title": "合成方法", "content": "仅为指引"})
        db.add(AnalysisRun(id="ANL-synthetic-report", case_id=case.id, status="COMPLETED", model_config_json=json_dumps(config),
                           result_json=json_dumps(payload(sample_markdown())), evidence_json=json_dumps(evidence)))
        db.commit()
    yield factory, managed
    engine.dispose()


def context():
    return report.get_report_context("CASE-synthetic-report", "ANL-synthetic-report")


def test_caller_transaction_not_closed_or_committed_and_pointer_resolved(isolated, monkeypatch):
    factory, _ = isolated
    with factory() as db:
        db.get(Case, "CASE-synthetic-report").description = "uncommitted"
        db.flush()
        def forbidden():
            raise AssertionError("must not commit or close caller transaction")
        with monkeypatch.context() as patch:
            patch.setattr(db, "commit", forbidden)
            patch.setattr(db, "close", forbidden)
            result = report.get_report_context("CASE-synthetic-report", "ANL-synthetic-report", db=db)
            assert result["case"].description == "uncommitted"
            assert result["report_template"]["version"] == 2
            assert result["case_evidence_ids"] == IDS
        db.rollback()


def test_new_snapshot_fixed_across_category_and_template_edits(isolated):
    factory, _ = isolated
    before = report_markdown(context())
    with factory() as db:
        db.get(Case, "CASE-synthetic-report").problem_category = "connection"
        db.add(KnowledgeDocument(id="DOC-new-template", title="新模板", content="NEW_TEMPLATE_MUST_NOT_APPEAR", active=True,
            review_status="ACTIVE", confidentiality="INTERNAL", metadata_json=json_dumps({"knowledge_role": "report_template", "problem_categories": ["network"]})))
        db.commit()
    assert report_markdown(context()) == before
    assert "固定问题类别：network" in before and " / v2" in before


def test_category_template_and_network_fallback_frozen_in_same_transaction(isolated):
    factory, _ = isolated
    with factory() as db:
        db.add(KnowledgeDocument(id="DOC-connection", title="连接规范", content="固定连接报告规范", active=True,
            review_status="ACTIVE", confidentiality="INTERNAL", metadata_json=json_dumps({"knowledge_role": "report_template", "problem_categories": ["connection"]})))
        db.flush()
        case = db.get(Case, "CASE-synthetic-report")
        case.problem_category = "connection"
        pinned = workbench.capture_configuration(db, case)
        assert workbench.resolve_configuration(db, pinned)["report_template"]["id"] == "DOC-connection"
        db.get(KnowledgeDocument, "DOC-connection").content = "CHANGED"
        case.problem_category = "unknown"
        assert workbench.resolve_configuration(db, pinned)["report_template"]["content"] == "固定连接报告规范"
        fallback = workbench.resolve_configuration(db, workbench.capture_configuration(db, case))
        assert fallback["report_template"]["id"] == "builtin-network-report"


def test_conservative_fallback_has_all_tables_without_invented_devices(isolated):
    ctx = context()
    ctx["result"]["report_markdown"] = ""
    rendered = report_markdown(ctx)
    assert [block["text"] for block in markdown_blocks(rendered) if block["kind"] == "heading" and block["level"] == 2] == list(CHAPTERS)
    assert len([block for block in markdown_blocks(rendered) if block["kind"] == "table"]) == 4
    assert "设备归属待确认" in rendered and "因果链待确认" in rendered
    assert "HG8145X6" not in rendered and "192.168.1.1" not in rendered
    assert "处理优先级：P0" in rendered and "置信度" in rendered


def test_coverage_deduplicates_by_method_and_never_counts_exclusions_as_roots():
    items = [dict(id="NODE1", method_document_id="METHOD1", status="SUPPORTED", evidence_ids=["EVT-synthetic-1"], attempted=True),
             dict(id="NODE2", method_document_id="METHOD1", status="EXCLUDED", evidence_ids=["EVT-synthetic-2"], attempted=True),
             dict(id="NODE3", method_document_id="METHOD1", status="SUPPORTED", evidence_ids=["DOC-guide"], attempted=True)]
    text = _coverage({"diagnostic_planning": {"fault_tree_coverage": {"items": items + [deepcopy(items[0])]}}}, IDS)
    assert "已定位根因：1/3" in text and "已排除：1/3" in text
    assert "账本缺少有效证据" in text


class Tags(HTMLParser):
    def __init__(self):
        super().__init__()
        self.tags, self.text = [], []
    def handle_starttag(self, tag, attrs):
        self.tags.append(tag)
    def handle_data(self, value):
        self.text.append(value)


def test_passive_html_and_valid_table_list_code_blocks():
    source = '# <script>alert(1)</script>\n\n| 列一 | 列二 |\n|---|---|\n| a\\|b | <img src=x onerror=bad> |\n\n- 中文列表\n- **加粗**\n\n3. 第三项\n\n```html\n<iframe src=x>\n```\n'
    rendered = markdown_html(source)
    parsed = Tags()
    parsed.feed(rendered)
    assert not {"script", "img", "iframe", "a"} & set(parsed.tags)
    assert {"thead", "tbody", "th", "td", "ul", "ol", "li", "pre", "code"} <= set(parsed.tags)
    assert parsed.tags.count("td") == 2
    assert "a|b" in rendered and '<ol start="3">' in rendered
    assert rendered.index("</table>") < rendered.index("<ul>") < rendered.index("<pre>")
    assert "&lt;script&gt;" in rendered and "&lt;img" in rendered


def test_all_formats_share_unicode_chapters_tables_lists_and_page_width(isolated):
    _, managed = isolated
    generated = [writer("CASE-synthetic-report", "ANL-synthetic-report") for writer in
                 (report.generate_markdown, report.generate_html_file, report.generate_docx, report.generate_pdf)]
    paths = {row.format: managed.resolve_path(row.stored_path) for row in generated}
    html = Tags()
    html.feed(paths["html"].read_text(encoding="utf-8"))
    document = Document(paths["docx"])
    word = "\n".join([p.text for p in document.paragraphs] + [cell.text for table in document.tables for row in table.rows for cell in row.cells])
    pdf = PdfReader(paths["pdf"])
    pdf_text = "".join(page.extract_text() for page in pdf.pages)
    for content in ("".join(html.text), word, re.sub(r"\s+", "", pdf_text)):
        for chapter in CHAPTERS:
            assert chapter in content
        assert "合成/ap1.log:L10-L12" in content
        assert "EVT-synthetic" not in content
        assert "处理优先级：P0" in content
    section = document.sections[0]
    available = section.page_width - section.left_margin - section.right_margin
    for table in document.tables:
        assert sum(column.width for column in table.columns) <= available + 20
        assert table.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None
    with ZipFile(paths["docx"]) as zipped:
        assert b'eastAsia="Microsoft YaHei"' in zipped.read("word/styles.xml")
    assert all(float(page.mediabox.width) > float(page.mediabox.height) for page in pdf.pages)
    assert len(document.tables) == 5
    assert html.tags.count("table") == 5
    for row in generated:
        assert sha256_file(paths[row.format]) == row.sha256


def test_legacy_files_and_analysis_remain_unchanged_when_new_export_created(isolated):
    factory, managed = isolated
    with factory() as db:
        analysis = db.get(AnalysisRun, "ANL-synthetic-report")
        analysis.model_config_json = "{}"
        analysis.result_json = json_dumps(payload("# 历史自由格式\n\n<script>bad</script>\n\n保留旧段落"))
        original = analysis.result_json
        db.commit()
    first = report.generate_html_file("CASE-synthetic-report", "ANL-synthetic-report")
    path = managed.resolve_path(first.stored_path)
    before = path.read_bytes()
    report.generate_html_file("CASE-synthetic-report", "ANL-synthetic-report")
    report.generate_docx("CASE-synthetic-report", "ANL-synthetic-report")
    report.generate_pdf("CASE-synthetic-report", "ANL-synthetic-report")
    assert path.read_bytes() == before
    assert b"<script>" not in before
    with factory() as db:
        assert db.get(AnalysisRun, "ANL-synthetic-report").result_json == original


def test_failed_export_discards_reservation_and_temporary_only(isolated, monkeypatch):
    factory, managed = isolated
    def fail(path):
        Path(path).write_text("partial", encoding="utf-8")
        raise RuntimeError("synthetic writer failure")
    with pytest.raises(RuntimeError, match="synthetic"):
        report._export("CASE-synthetic-report", "ANL-synthetic-report", "docx", fail)
    with factory() as db:
        assert not list(db.scalars(select(Report)))
    assert not list(managed.report_dir("CASE-synthetic-report").iterdir())


def test_every_ap_requires_analysis_and_conclusion_rows():
    text = sample_markdown()
    start = text.index("### AP2 — 离线", text.index("## 三、"))
    with pytest.raises(ValueError, match="Every identified"):
        validate_report_markdown(text[:start] + text[text.index("## 四、"):], IDS, builtin_report_template())
    text = re.sub(r"\| AP2 \| 离线 \| 链路中断[^\n]+\n", "", text)
    with pytest.raises(ValueError, match="root-cause overview"):
        validate_report_markdown(text, IDS, builtin_report_template())


def test_confidence_definition_and_table_width_are_validated():
    for text in (sample_markdown().replace("P3（可能）", "P1（可能）"),
                 sample_markdown().replace("| AP1 | 待确认", "| AP1 | 额外列 | 待确认", 1)):
        with pytest.raises(ValueError):
            validate_report_markdown(text, IDS, builtin_report_template())


def test_malicious_filename_does_not_change_tables_or_create_active_html(isolated):
    ctx = context()
    ctx["evidence_labels"]["EVT-synthetic-1"] = '<img src=x onerror=bad>| forged\n## injected:L1'
    rendered = report.html_report(ctx)
    parsed = Tags()
    parsed.feed(rendered)
    assert "img" not in parsed.tags
    assert "&lt;img" in rendered and "| forged" in rendered
    assert parsed.tags.count("h2") == 4 and parsed.tags.count("table") == 5


def test_missing_or_tampered_rc_snapshot_fails_before_report_publication(isolated):
    factory, _ = isolated
    with factory() as db:
        row = db.scalars(select(WorkbenchRecord).where(WorkbenchRecord.kind == "run_context")).one()
        row.payload_json = "{}"
        db.commit()
    with pytest.raises(ValueError, match="快照"):
        report.generate_html_file("CASE-synthetic-report", "ANL-synthetic-report")
    with factory() as db:
        assert not list(db.scalars(select(Report)))


def test_unicode_symbols_embed_actual_glyphs_in_pdf(isolated):
    factory, managed = isolated
    with factory() as db:
        db.get(AnalysisRun, "ANL-synthetic-report").result_json = json_dumps(payload(sample_markdown().replace("异常 ❌", "异常 **❌**")))
        db.commit()
    row = report.generate_pdf("CASE-synthetic-report", "ANL-synthetic-report")
    reader = PdfReader(managed.resolve_path(row.stored_path))
    fonts = [font.get_object() for page in reader.pages for font in page["/Resources"]["/Font"].values()]
    assert any("SegoeUISymbol" in str(font.get("/BaseFont")) for font in fonts)
    assert pdfmetrics.getFont("WorkbenchSymbols").face.charToGlyph.get(ord("❌"))
    assert pdfmetrics.getFont("WorkbenchCJK-Bold").face.charToGlyph.get(ord("链"))
    assert "❌" in "".join(page.extract_text() for page in reader.pages)


def test_large_chinese_table_rows_paginate_without_dropping_content(isolated):
    factory, managed = isolated
    with factory() as db:
        analysis = db.get(AnalysisRun, "ANL-synthetic-report")
        analysis.model_config_json = json_dumps({"problem_category": "connection", "report_template": template("请保留观测记录和证据。")})
        data = payload("# 长表格合成检查\n\n| 观察 | 证据 |\n|---|---|\n| " + "中文观测；" * 1500 + "末尾标识 | [[EVT-synthetic-1]] |\n")
        analysis.result_json = json_dumps(data)
        db.commit()
    row = report.generate_pdf("CASE-synthetic-report", "ANL-synthetic-report")
    pdf = PdfReader(managed.resolve_path(row.stored_path))
    assert len(pdf.pages) > 1
    assert "末尾标识" in "".join(page.extract_text() for page in pdf.pages)


def test_new_report_needs_case_allowlist_and_unresolved_pointer_is_rejected():
    with pytest.raises(ValueError, match="case_evidence_ids"):
        validate_llm_diagnosis(payload(sample_markdown()), IDS | {"DOC-guide"}, report_template=builtin_report_template())
    with pytest.raises(ValueError, match="Resolve workbench"):
        report_instructions({"workbench_snapshot_id": "RC-unresolved"})


def test_escaped_pipe_backslash_and_degenerate_table_lines():
    assert table_cells(r"| a\|b | c |") == [r"a\|b", "c"]
    assert table_cells(r"| a\\| b |") == ["a\\\\", "b"]
    assert not any(block["kind"] == "table" for block in markdown_blocks("|\n|---|\n|\n|\n"))


def test_suggestion_and_cross_category_notes_cannot_inject_headings(isolated):
    ctx = context()
    ctx["result"].update(suggested_problem_category="WB-SYNTHETIC", category_reason="参考观察\n## 伪造标题 [[EVT-synthetic-1]]")
    ctx["evidence"]["DOC-guide"]["metadata"] = {"cross_category_reason": "该类别暂无命中，核对适用性\n## forged", "problem_categories": ["connection"]}
    text = report_markdown(ctx)
    assert "来源类别：connection" in text and "该类别暂无命中" in text
    assert "DOC-guide" not in text and "EVT-synthetic-1" not in text
    assert len([block for block in markdown_blocks(text) if block["kind"] == "heading" and block["level"] == 2]) == 4
