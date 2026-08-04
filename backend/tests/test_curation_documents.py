from pathlib import Path

from docx import Document
import pytest
from reportlab.pdfgen import canvas

from app.services.curation_documents import (
    DocumentExtractionError,
    prepare_curation_document,
)


def test_html_extracts_visible_text_and_ignores_active_or_hidden_content(
    tmp_path: Path,
) -> None:
    source = tmp_path / "case.html"
    source.write_text(
        """<!doctype html>
<html><head><title>AP failure</title><style>.x { display:none }</style></head>
<body>
  <h1>Authentication timeout</h1>
  <script>IGNORE_INJECTION()</script>
  <p hidden>hidden secret</p>
  <table><tr><th>Cause</th><th>Solution</th></tr>
  <tr><td>Key mismatch</td><td>Replace key</td></tr></table>
</body></html>
""",
        encoding="utf-8",
    )
    result = prepare_curation_document(source, "case.html", tmp_path / "html.txt")

    assert result is not None
    assert result.method == "html_visible_text"
    extracted = result.path.read_text(encoding="utf-8")
    assert "Authentication timeout" in extracted
    assert "Key mismatch" in extracted
    assert "Replace key" in extracted
    assert "IGNORE_INJECTION" not in extracted
    assert "hidden secret" not in extracted


def test_docx_extracts_paragraphs_and_tables(tmp_path: Path) -> None:
    source = tmp_path / "case.docx"
    document = Document()
    document.add_heading("AP association failure", level=1)
    document.add_paragraph("Authentication failed after key exchange")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Cause"
    table.cell(0, 1).text = "Solution"
    table.cell(1, 0).text = "Key mismatch"
    table.cell(1, 1).text = "Replace key"
    document.save(source)

    result = prepare_curation_document(source, "case.docx", tmp_path / "docx.txt")

    assert result is not None
    assert result.method == "docx_paragraphs_tables"
    extracted = result.path.read_text(encoding="utf-8")
    assert "AP association failure" in extracted
    assert "Authentication failed after key exchange" in extracted
    assert "Cause | Solution" in extracted
    assert "Key mismatch | Replace key" in extracted


def test_pdf_extracts_text_layer_with_page_markers(tmp_path: Path) -> None:
    source = tmp_path / "case.pdf"
    pdf = canvas.Canvas(str(source))
    pdf.drawString(72, 760, "AP authentication timeout")
    pdf.drawString(72, 740, "Cause: shared key mismatch")
    pdf.showPage()
    pdf.drawString(72, 760, "Solution: replace key and restart WLAN")
    pdf.save()

    result = prepare_curation_document(source, "case.pdf", tmp_path / "pdf.txt")

    assert result is not None
    assert result.method == "pdf_text_layer"
    assert result.page_count == 2
    extracted = result.path.read_text(encoding="utf-8")
    assert "[PDF page 1]" in extracted
    assert "AP authentication timeout" in extracted
    assert "[PDF page 2]" in extracted
    assert "Solution: replace key and restart WLAN" in extracted


def test_legacy_doc_requires_explicit_conversion(tmp_path: Path) -> None:
    source = tmp_path / "legacy.doc"
    source.write_bytes(b"legacy-word-binary")

    with pytest.raises(DocumentExtractionError) as captured:
        prepare_curation_document(source, "legacy.doc", tmp_path / "legacy.txt")

    assert captured.value.code == "legacy_doc_requires_conversion"


def test_pdf_without_a_text_layer_is_rejected_as_non_extractable(tmp_path: Path) -> None:
    source = tmp_path / "scan.pdf"
    pdf = canvas.Canvas(str(source))
    pdf.showPage()
    pdf.save()

    with pytest.raises(DocumentExtractionError) as captured:
        prepare_curation_document(source, "scan.pdf", tmp_path / "scan.txt")

    assert captured.value.code == "document_has_no_extractable_text"
