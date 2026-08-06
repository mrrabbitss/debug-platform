"""Generate deterministic synthetic DOCX/PDF fixtures for the golden corpus.

This script must be run with the bundled Codex document runtime (or a Python
environment containing python-docx and reportlab). It never reads runtime or
company data; all fixture text is defined below.
"""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import shutil
import sys
import tempfile
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab import rl_config
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


REPO_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = REPO_ROOT / "sample_data" / "golden_incident" / "case"
CORPUS_PATH = OUTPUT_DIR.parent / "corpus.json"
FIXED_ZIP_TIME = (2026, 3, 2, 3, 29, 16)
FIXED_CREATED = datetime(2026, 3, 2, 3, 29, 16, tzinfo=timezone.utc)
GENERATED_FIXTURES = ("solution.docx", "validation.pdf")


def _set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _normalize_docx_zip(path: Path) -> None:
    """Rewrite OOXML entries with stable order, timestamps and permissions."""
    with tempfile.TemporaryDirectory(prefix="golden-docx-") as temporary:
        normalized = Path(temporary) / path.name
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
            normalized,
            "w",
            compression=zipfile.ZIP_DEFLATED,
            compresslevel=9,
        ) as destination:
            for name in sorted(source.namelist()):
                payload = source.read(name)
                info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.create_system = 0
                info.external_attr = 0
                destination.writestr(info, payload)
        shutil.copyfile(normalized, path)


def generate_docx(path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(9.5)
    styles["Title"].font.name = "Arial"
    styles["Title"].font.size = Pt(20)
    styles["Title"].font.color.rgb = RGBColor(24, 58, 92)
    styles["Heading 1"].font.name = "Arial"
    styles["Heading 1"].font.size = Pt(12)
    styles["Heading 1"].font.color.rgb = RGBColor(24, 58, 92)

    header = section.header.paragraphs[0]
    header.text = "GW/AP DEBUG PLATFORM  |  SYNTHETIC GOLDEN INCIDENT"
    header.style = styles["Caption"]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    title = document.add_paragraph("Recovery Procedure", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle = document.add_paragraph("Golden fixture · WLAN authentication · 2026-03-02")
    subtitle.runs[0].font.color.rgb = RGBColor(91, 107, 122)

    table = document.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    details = [
        ("Fixture", "GOLDEN-AP-01"),
        ("Trigger", "AUTH_TIMEOUT after the 4-way handshake begins"),
        ("Verified cause", "Shared-key mismatch in the lab WLAN profile"),
        ("Safety", "Synthetic data; no production identifiers"),
    ]
    for row, (label, value) in zip(table.rows, details, strict=True):
        row.cells[0].text = label
        row.cells[1].text = value
        _set_cell_shading(row.cells[0], "E8F0F7")
        row.cells[0].paragraphs[0].runs[0].bold = True

    document.add_paragraph("Approved steps", style="Heading 1")
    steps = [
        "Back up the current WLAN configuration before making a change.",
        "Replace the synthetic lab shared key with the approved matching value.",
        "Restart only the WLAN authentication service; do not reboot the device.",
        "Associate the synthetic client and verify that AUTH_TIMEOUT no longer appears.",
    ]
    for step in steps:
        document.add_paragraph(step, style="List Number")

    document.add_paragraph("Rollback and limits", style="Heading 1")
    document.add_paragraph(
        "If association still fails, restore the backup and stop. This fixture does not "
        "justify replacing hardware or declaring a firmware defect."
    )
    document.core_properties.title = "Synthetic Golden Incident Recovery Procedure"
    document.core_properties.subject = "Deterministic E2E and curation fixture"
    document.core_properties.author = "GW/AP Debug Platform"
    document.core_properties.created = FIXED_CREATED
    document.core_properties.modified = FIXED_CREATED
    document.save(path)
    _normalize_docx_zip(path)


def generate_pdf(path: Path) -> None:
    rl_config.invariant = 1
    styles = getSampleStyleSheet()
    title = ParagraphStyle(
        "GoldenTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        textColor=HexColor("#183A5C"),
        alignment=TA_LEFT,
        spaceAfter=8,
    )
    heading = ParagraphStyle(
        "GoldenHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=15,
        textColor=HexColor("#183A5C"),
        spaceBefore=8,
        spaceAfter=5,
    )
    body = ParagraphStyle(
        "GoldenBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=13,
        textColor=HexColor("#263746"),
        spaceAfter=5,
    )
    document = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=17 * mm,
        bottomMargin=17 * mm,
        title="Synthetic Golden Incident Validation",
        author="GW/AP Debug Platform",
        creator="GW/AP Debug Platform golden fixture generator",
        invariant=1,
    )
    story = [
        Paragraph("Validation Record", title),
        Paragraph("GOLDEN-AP-01 · WLAN authentication · synthetic fixture", body),
        Spacer(1, 3 * mm),
        Table(
            [
                ["Check", "Expected observation", "Result"],
                ["Association", "Client completes authentication", "PASS"],
                ["Log", "No new AUTH_TIMEOUT event", "PASS"],
                ["Configuration", "Shared keys match", "PASS"],
            ],
            colWidths=[36 * mm, 100 * mm, 24 * mm],
            style=TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), HexColor("#183A5C")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), HexColor("#FFFFFF")),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("GRID", (0, 0), (-1, -1), 0.5, HexColor("#A9B8C6")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [HexColor("#F7FAFC"), HexColor("#FFFFFF")]),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ]
            ),
        ),
        Paragraph("Conclusion", heading),
        Paragraph(
            "Authentication succeeds after the approved shared-key update. The lab "
            "observation window contains no recurrence. This validates only the synthetic "
            "fixture and must not be generalized to a customer environment without evidence.",
            body,
        ),
        Paragraph("Evidence boundary", heading),
        Paragraph(
            "The result supports the shared-key mismatch diagnosis. It does not support a "
            "firmware defect, hardware replacement, or a claim of customer impact.",
            body,
        ),
    ]
    document.build(story)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _generate_documents(output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    generate_docx(output_dir / "solution.docx")
    generate_pdf(output_dir / "validation.pdf")


def _check_documents() -> int:
    manifest = json.loads(CORPUS_PATH.read_text(encoding="utf-8"))
    expected_hashes = manifest.get("fixture_sha256", {})
    failures: list[str] = []

    with tempfile.TemporaryDirectory(prefix="golden-fixture-check-") as temporary:
        generated_dir = Path(temporary)
        _generate_documents(generated_dir)
        for filename in GENERATED_FIXTURES:
            relative_path = f"case/{filename}"
            expected = expected_hashes.get(relative_path)
            committed_path = OUTPUT_DIR / filename
            if not expected:
                failures.append(f"missing fixture_sha256 entry: {relative_path}")
                continue
            if not committed_path.is_file():
                failures.append(f"missing committed fixture: {relative_path}")
                continue

            committed_hash = _sha256(committed_path)
            generated_hash = _sha256(generated_dir / filename)
            if committed_hash != expected:
                failures.append(
                    f"committed {relative_path} hash {committed_hash} != manifest {expected}"
                )
            if generated_hash != expected:
                failures.append(
                    f"generated {relative_path} hash {generated_hash} != manifest {expected}"
                )

    if failures:
        for failure in failures:
            print(f"ERROR: {failure}", file=sys.stderr)
        return 1
    print(f"Verified {len(GENERATED_FIXTURES)} deterministic golden documents")
    return 0


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--check",
        action="store_true",
        help="regenerate in a temporary directory and verify committed manifest hashes",
    )
    return parser.parse_args()


def main() -> int:
    args = _parse_args()
    if args.check:
        return _check_documents()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    _generate_documents(OUTPUT_DIR)
    print(f"Generated golden documents in {OUTPUT_DIR}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
